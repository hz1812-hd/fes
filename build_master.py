# -*- coding: utf-8 -*-
"""
CASE 3 - FINREPORT INTELLIGENCE
BUOC 1 (+2): HOP NHAT DU LIEU THANH FILE MASTER BANG PYTHON
             (thay cho phan lam thu cong tren Excel)

Dau vao (Data Pack, thu muc 01_data):
    company_info.csv        - 1 dong / doanh nghiep           (khoa: company_id)
    income_statement.csv    - 1 dong / doanh nghiep-nam       (khoa: company_id + fiscal_year)
    balance_sheet.csv       - nt
    cashflow.csv            - nt
    financial_notes.csv     - nt

Dau ra (thu muc output_master/):
    1. master_hopnhat.xlsx               - bang master 60 dong (bo cuc giong file Excel cua ban:
                                           dong 1-2 tieu de, dong 4 ten cot, du lieu tu dong 5)
    2. master_hopnhat_with_metrics.xlsx  - master + 17 chi tieu tai chinh (README muc 6)
    3. master_hopnhat.csv                - ban CSV (tieu de o dong 1) de dung cho cac buoc sau
    4. master_hopnhat_with_metrics.csv
    5. bao_cao_kiem_tra.txt              - bao cao kiem tra chat luong du lieu

Cach chay trong VS Code (terminal):
    py -m pip install pandas numpy openpyxl
    py build_master.py --data-dir "duong/dan/CASE3_FINREPORT_INTELLIGENCE/01_data"

    # neu cac file csv nam ngay trong thu muc hien tai thi chi can:
    py build_master.py

File dau ra "master_hopnhat_with_metrics.xlsx" tuong thich truc tiep voi benchmark_nganh.py:
    py benchmark_nganh.py --input output_master/master_hopnhat_with_metrics.xlsx

Quy uoc (README muc 7): nam 2021 khong co du lieu 2020 -> cac chi tieu can nam truoc
de TRONG (NaN). KHONG tu sinh so lieu 2020.
"""

import argparse
import os
import sys

import numpy as np
import pandas as pd

# ----------------------------------------------------------------------------
# 1. CAU HINH
# ----------------------------------------------------------------------------
KEYS = ["company_id", "fiscal_year"]
TOL = 0.5  # dung sai khi doi chieu (don vi: trieu VND; so lieu la so nguyen)

FILE_COMPANY = "company_info.csv"
FILES_YEARLY = ["income_statement.csv", "balance_sheet.csv", "cashflow.csv", "financial_notes.csv"]

# Thu tu cot cua file master (giong file Excel master ban da lam).
# Cot nao khong co trong du lieu se bi bo qua; cot la se duoc dat o cuoi.
MASTER_ORDER = [
    'company_id', 'fiscal_year', 'revenue_vnd_mn', 'cost_of_goods_sold_vnd_mn',
    'gross_profit_vnd_mn', 'selling_expense_vnd_mn', 'administrative_expense_vnd_mn',
    'depreciation_amortization_vnd_mn', 'operating_profit_vnd_mn', 'finance_income_vnd_mn',
    'interest_expense_vnd_mn', 'other_income_vnd_mn', 'profit_before_tax_vnd_mn',
    'income_tax_expense_vnd_mn', 'net_income_vnd_mn', 'cash_and_cash_equivalents_vnd_mn',
    'accounts_receivable_vnd_mn', 'inventory_vnd_mn', 'other_current_assets_vnd_mn',
    'total_current_assets_vnd_mn', 'property_plant_equipment_net_vnd_mn',
    'intangible_assets_vnd_mn', 'long_term_investments_vnd_mn',
    'other_noncurrent_assets_vnd_mn', 'total_assets_vnd_mn', 'short_term_debt_vnd_mn',
    'accounts_payable_vnd_mn', 'other_current_liabilities_vnd_mn',
    'total_current_liabilities_vnd_mn', 'long_term_debt_vnd_mn',
    'other_noncurrent_liabilities_vnd_mn', 'total_liabilities_vnd_mn',
    'share_capital_vnd_mn', 'retained_earnings_vnd_mn', 'other_equity_vnd_mn',
    'total_equity_vnd_mn', 'total_liabilities_and_equity_vnd_mn',
    'cash_effect_change_accounts_receivable_vnd_mn', 'cash_effect_change_inventory_vnd_mn',
    'cash_effect_change_other_current_assets_vnd_mn',
    'cash_effect_change_accounts_payable_vnd_mn',
    'cash_effect_change_other_current_liabilities_vnd_mn',
    'cash_effect_change_other_noncurrent_liabilities_vnd_mn',
    'net_cash_from_operating_activities_vnd_mn', 'capital_expenditure_vnd_mn',
    'change_intangible_assets_vnd_mn', 'change_long_term_investments_vnd_mn',
    'change_other_noncurrent_assets_vnd_mn', 'net_cash_from_investing_activities_vnd_mn',
    'net_debt_issuance_repayment_vnd_mn', 'equity_issuance_vnd_mn', 'dividends_paid_vnd_mn',
    'net_cash_from_financing_activities_vnd_mn', 'net_change_in_cash_vnd_mn',
    'beginning_cash_vnd_mn', 'ending_cash_vnd_mn', 'average_customer_credit_terms_days',
    'inventory_policy_days_target', 'average_borrowing_rate_pct',
    'debt_maturing_within_12m_pct', 'major_customer_concentration_pct',
    'export_revenue_pct', 'capex_intensity_pct', 'auditor_opinion_simulated', 'note_text',
    'company_name', 'ticker_simulated', 'industry', 'subindustry', 'headquarters',
    'founded_year', 'listing_status_simulated', 'fiscal_year_end', 'reporting_currency',
    'statement_unit', 'company_size_band', 'data_is_synthetic',
    'check_1_assets_eq_liab_equity', 'check_2_cash_reconciliation',
    'check_3_ending_cash_eq_bs_cash',
]

METRIC_ORDER = [
    "revenue_growth_pct", "gross_margin_pct", "operating_margin_pct", "net_margin_pct",
    "current_ratio", "quick_ratio", "debt_to_assets_pct", "debt_to_equity",
    "roa_pct", "roe_pct", "cfo_to_net_income", "free_cash_flow_vnd_mn",
    "dso_days", "inventory_days", "dpo_days", "interest_coverage", "working_capital_vnd_mn",
]

REPORT = []  # cac dong cua bao cao kiem tra


def log(msg=""):
    print(msg)
    REPORT.append(msg)


def fatal(msg):
    log(f"[LOI] {msg}")
    save_report_on_error()
    sys.exit(1)


_REPORT_PATH = None


def save_report_on_error():
    if _REPORT_PATH:
        os.makedirs(os.path.dirname(_REPORT_PATH), exist_ok=True)
        with open(_REPORT_PATH, "w", encoding="utf-8") as f:
            f.write("\n".join(REPORT))


# ----------------------------------------------------------------------------
# 2. DOC DU LIEU
# ----------------------------------------------------------------------------
def find_file(data_dir: str, name: str) -> str:
    """Tim file theo ten: thu muc chi dinh, thu muc con 01_data, thu muc hien tai, roi quet de quy."""
    candidates = [
        os.path.join(data_dir, name),
        os.path.join(data_dir, "01_data", name),
        os.path.join(".", name),
        os.path.join(".", "01_data", name),
    ]
    for p in candidates:
        if os.path.isfile(p):
            return p
    for root, _, files in os.walk(data_dir):
        if name in files:
            return os.path.join(root, name)
    fatal(f"Khong tim thay file '{name}' (da tim trong: {data_dir}). "
          f"Hay truyen dung duong dan bang --data-dir.")


def read_csv(path: str) -> pd.DataFrame:
    df = pd.read_csv(path, encoding="utf-8-sig")
    df.columns = [c.strip() for c in df.columns]
    return df


def check_keys(df: pd.DataFrame, name: str, keys):
    miss = [k for k in keys if k not in df.columns]
    if miss:
        fatal(f"{name}: thieu cot khoa {miss}")
    if df.duplicated(keys).any():
        fatal(f"{name}: co khoa {keys} bi trung lap")
    if df[keys].isna().any().any():
        fatal(f"{name}: co khoa bi de trong")


# ----------------------------------------------------------------------------
# 3. HOP NHAT
# ----------------------------------------------------------------------------
def merge_without_dups(base: pd.DataFrame, other: pd.DataFrame, keys, name: str, how: str,
                       validate: str) -> pd.DataFrame:
    """Merge; cot trung ten (ngoai khoa) duoc doi chieu roi chi giu 1 ban."""
    overlap = [c for c in other.columns if c in base.columns and c not in keys]
    if overlap:
        tmp = base[keys + overlap].merge(other[keys + overlap], on=keys, suffixes=("", "__o"))
        for c in overlap:
            a, b = tmp[c], tmp[c + "__o"]
            diff = ~((a == b) | (a.isna() & b.isna()))
            if diff.any():
                log(f"[CANH BAO] {name}: cot trung '{c}' co {int(diff.sum())} dong khac nhau "
                    f"-> giu gia tri tu file dau tien")
        other = other.drop(columns=overlap)
        log(f"   - {name}: bo {len(overlap)} cot trung ten da co san ({', '.join(overlap)})")
    merged = base.merge(other, on=keys, how=how, validate=validate, indicator="_m")
    unmatched = merged["_m"] != "both"
    if unmatched.any():
        fatal(f"{name}: co {int(unmatched.sum())} dong khong khop khoa voi cac file truoc")
    return merged.drop(columns="_m")


def build_master(data_dir: str) -> pd.DataFrame:
    log("=" * 70)
    log("BUOC 1. DOC VA HOP NHAT DU LIEU")
    log("=" * 70)

    yearly = {}
    for f in FILES_YEARLY:
        path = find_file(data_dir, f)
        df = read_csv(path)
        check_keys(df, f, KEYS)
        yearly[f] = df
        log(f"[OK] {f:<24} {df.shape[0]:>3} dong x {df.shape[1]:>2} cot  ({path})")

    cpath = find_file(data_dir, FILE_COMPANY)
    comp = read_csv(cpath)
    check_keys(comp, FILE_COMPANY, ["company_id"])
    log(f"[OK] {FILE_COMPANY:<24} {comp.shape[0]:>3} dong x {comp.shape[1]:>2} cot  ({cpath})")

    # Kiem tra so dong dong nhat giua cac file company-year
    sizes = {k: len(v) for k, v in yearly.items()}
    if len(set(sizes.values())) != 1:
        log(f"[CANH BAO] So dong cac file khong dong nhat: {sizes}")

    log("\nHop nhat theo khoa company_id + fiscal_year:")
    master = yearly[FILES_YEARLY[0]].copy()
    for f in FILES_YEARLY[1:]:
        master = merge_without_dups(master, yearly[f], KEYS, f, how="outer", validate="one_to_one")

    log("Hop nhat thong tin doanh nghiep theo company_id:")
    unknown = set(master["company_id"]) - set(comp["company_id"])
    if unknown:
        fatal(f"company_id khong co trong company_info.csv: {sorted(unknown)}")
    master = merge_without_dups(master, comp, ["company_id"], FILE_COMPANY, how="left",
                                validate="many_to_one")

    master["fiscal_year"] = master["fiscal_year"].astype(int)
    master = master.sort_values(KEYS).reset_index(drop=True)
    return master


# ----------------------------------------------------------------------------
# 4. KIEM TRA DOI CHIEU (README muc 3)
# ----------------------------------------------------------------------------
def need(df, cols, label):
    miss = [c for c in cols if c not in df.columns]
    if miss:
        log(f"[BO QUA] {label}: thieu cot {miss}")
        return False
    return True


def add_checks(df: pd.DataFrame) -> pd.DataFrame:
    log("\n" + "=" * 70)
    log("BUOC 2. KIEM TRA DOI CHIEU GIUA CAC BAO CAO")
    log("=" * 70)
    n = len(df)

    if need(df, ["total_assets_vnd_mn", "total_liabilities_vnd_mn", "total_equity_vnd_mn"], "check_1"):
        df["check_1_assets_eq_liab_equity"] = (
            (df["total_assets_vnd_mn"] - df["total_liabilities_vnd_mn"] - df["total_equity_vnd_mn"]).abs() <= TOL)
    if need(df, ["beginning_cash_vnd_mn", "net_change_in_cash_vnd_mn", "ending_cash_vnd_mn"], "check_2"):
        df["check_2_cash_reconciliation"] = (
            (df["beginning_cash_vnd_mn"] + df["net_change_in_cash_vnd_mn"] - df["ending_cash_vnd_mn"]).abs() <= TOL)
    if need(df, ["ending_cash_vnd_mn", "cash_and_cash_equivalents_vnd_mn"], "check_3"):
        df["check_3_ending_cash_eq_bs_cash"] = (
            (df["ending_cash_vnd_mn"] - df["cash_and_cash_equivalents_vnd_mn"]).abs() <= TOL)

    labels = {
        "check_1_assets_eq_liab_equity": "Total Assets = Total Liabilities + Total Equity",
        "check_2_cash_reconciliation": "Beginning Cash + Net Change in Cash = Ending Cash",
        "check_3_ending_cash_eq_bs_cash": "Ending Cash = Cash tren balance_sheet",
    }
    for col, text in labels.items():
        if col in df.columns:
            ok = int(df[col].sum())
            flag = "DAT" if ok == n else "CO DONG LOI"
            log(f"[{flag}] {text}: {ok}/{n} dong dung")
            if ok != n:
                bad = df.loc[~df[col], KEYS].values.tolist()
                log(f"        Dong sai: {bad}")

    # Kiem tra bo sung (chi bao cao, khong ghi vao master)
    if need(df, ["revenue_vnd_mn", "cost_of_goods_sold_vnd_mn", "gross_profit_vnd_mn"], "gross profit"):
        bad = ((df["revenue_vnd_mn"] - df["cost_of_goods_sold_vnd_mn"] - df["gross_profit_vnd_mn"]).abs() > TOL).sum()
        log(f"[{'DAT' if bad == 0 else 'CANH BAO'}] Gross Profit = Revenue - COGS: {n - int(bad)}/{n} dong dung")
    if need(df, ["beginning_cash_vnd_mn", "ending_cash_vnd_mn"], "lien tuc tien"):
        prev_end = df.groupby("company_id")["ending_cash_vnd_mn"].shift(1)
        mask = prev_end.notna()
        bad = ((df.loc[mask, "beginning_cash_vnd_mn"] - prev_end[mask]).abs() > TOL).sum()
        log(f"[{'DAT' if bad == 0 else 'CANH BAO'}] Beginning Cash nam t = Ending Cash nam t-1: "
            f"{int(mask.sum()) - int(bad)}/{int(mask.sum())} cap nam dung")
    if "statement_unit" in df.columns:
        units = df["statement_unit"].unique().tolist()
        log(f"[{'DAT' if units == ['VND_million'] else 'CANH BAO'}] Don vi so lieu: {units}")
    return df


def quality_summary(df: pd.DataFrame):
    log("\n" + "=" * 70)
    log("TOM TAT CHAT LUONG DU LIEU")
    log("=" * 70)
    log(f"So dong: {len(df)} | So doanh nghiep: {df['company_id'].nunique()} | "
        f"Nam: {df['fiscal_year'].min()}-{df['fiscal_year'].max()} | So cot: {df.shape[1]}")
    per_company = df.groupby("company_id").size()
    if per_company.nunique() != 1:
        log(f"[CANH BAO] So nam moi doanh nghiep khong dong nhat: {per_company.to_dict()}")
    na = df.isna().sum()
    na = na[na > 0]
    if len(na):
        log(f"[CANH BAO] Cot co o trong: {na.to_dict()}")
    else:
        log("[DAT] Khong co o trong trong bang master")


# ----------------------------------------------------------------------------
# 5. TINH 17 CHI TIEU TAI CHINH (README muc 6)
# ----------------------------------------------------------------------------
def add_metrics(df: pd.DataFrame) -> pd.DataFrame:
    log("\n" + "=" * 70)
    log("BUOC 3. TINH 17 CHI TIEU TAI CHINH")
    log("=" * 70)
    required = [
        "revenue_vnd_mn", "cost_of_goods_sold_vnd_mn", "gross_profit_vnd_mn", "operating_profit_vnd_mn",
        "interest_expense_vnd_mn", "net_income_vnd_mn", "cash_and_cash_equivalents_vnd_mn",
        "accounts_receivable_vnd_mn", "inventory_vnd_mn", "total_current_assets_vnd_mn",
        "total_assets_vnd_mn", "short_term_debt_vnd_mn", "accounts_payable_vnd_mn",
        "total_current_liabilities_vnd_mn", "long_term_debt_vnd_mn", "total_equity_vnd_mn",
        "net_cash_from_operating_activities_vnd_mn", "capital_expenditure_vnd_mn",
    ]
    miss = [c for c in required if c not in df.columns]
    if miss:
        fatal(f"Khong the tinh chi tieu, thieu cot: {miss}")

    g = df.groupby("company_id")

    def prev(c):          # gia tri nam truoc cua cung doanh nghiep (NaN voi nam dau)
        return g[c].shift(1)

    def avg(c):           # so du binh quan dau-cuoi nam
        return (df[c] + prev(c)) / 2

    def div(a, b):        # chia an toan: mau = 0 -> NaN
        return a / b.replace(0, np.nan)

    m = pd.DataFrame(index=df.index)
    debt = df["short_term_debt_vnd_mn"] + df["long_term_debt_vnd_mn"]

    m["revenue_growth_pct"] = (df["revenue_vnd_mn"] - prev("revenue_vnd_mn")) / prev("revenue_vnd_mn") * 100
    m["gross_margin_pct"] = div(df["gross_profit_vnd_mn"], df["revenue_vnd_mn"]) * 100
    m["operating_margin_pct"] = div(df["operating_profit_vnd_mn"], df["revenue_vnd_mn"]) * 100
    m["net_margin_pct"] = div(df["net_income_vnd_mn"], df["revenue_vnd_mn"]) * 100
    m["current_ratio"] = div(df["total_current_assets_vnd_mn"], df["total_current_liabilities_vnd_mn"])
    m["quick_ratio"] = div(df["cash_and_cash_equivalents_vnd_mn"] + df["accounts_receivable_vnd_mn"],
                           df["total_current_liabilities_vnd_mn"])
    m["debt_to_assets_pct"] = div(debt, df["total_assets_vnd_mn"]) * 100
    m["debt_to_equity"] = div(debt, df["total_equity_vnd_mn"])
    m["roa_pct"] = div(df["net_income_vnd_mn"], avg("total_assets_vnd_mn")) * 100
    m["roe_pct"] = div(df["net_income_vnd_mn"], avg("total_equity_vnd_mn")) * 100
    m["cfo_to_net_income"] = div(df["net_cash_from_operating_activities_vnd_mn"], df["net_income_vnd_mn"])
    m["free_cash_flow_vnd_mn"] = (df["net_cash_from_operating_activities_vnd_mn"]
                                  + df["capital_expenditure_vnd_mn"])   # CAPEX da mang dau am
    m["dso_days"] = div(avg("accounts_receivable_vnd_mn"), df["revenue_vnd_mn"]) * 365
    m["inventory_days"] = div(avg("inventory_vnd_mn"), df["cost_of_goods_sold_vnd_mn"]) * 365
    m["dpo_days"] = div(avg("accounts_payable_vnd_mn"), df["cost_of_goods_sold_vnd_mn"]) * 365
    m["interest_coverage"] = div(df["operating_profit_vnd_mn"], df["interest_expense_vnd_mn"])
    m["working_capital_vnd_mn"] = df["total_current_assets_vnd_mn"] - df["total_current_liabilities_vnd_mn"]

    m = m[METRIC_ORDER]
    n21 = m[df["fiscal_year"] == df["fiscal_year"].min()].isna().sum()
    blanks = n21[n21 > 0].index.tolist()
    log(f"[OK] Da tinh {m.shape[1]} chi tieu cho {len(df)} dong")
    log(f"     Nam dau ({df['fiscal_year'].min()}) de trong vi can nam truoc: {blanks}")
    neg = df.loc[df["net_income_vnd_mn"] <= 0, KEYS].values.tolist()
    if neg:
        log(f"[LUU Y] Dong co loi nhuan rong <= 0 (CFO/Net Income kho dien giai): {neg}")
    return pd.concat([df, m], axis=1)


# ----------------------------------------------------------------------------
# 6. XUAT FILE
# ----------------------------------------------------------------------------
def order_columns(df: pd.DataFrame) -> pd.DataFrame:
    first = [c for c in MASTER_ORDER if c in df.columns]
    metric = [c for c in METRIC_ORDER if c in df.columns]
    rest = [c for c in df.columns if c not in first and c not in metric]
    # cot la (khong co trong master mau) dat truoc cot chi tieu
    return df[first + rest + metric]


def write_xlsx(df: pd.DataFrame, path: str):
    from openpyxl.styles import Alignment, Font, PatternFill
    n_rows, n_cols = df.shape
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Sheet1", index=False, startrow=3)   # tieu de o dong 4
        ws = xw.sheets["Sheet1"]
        ws["A1"] = f"BẢNG DỮ LIỆU MASTER SAU HỢP NHẤT ({n_rows} DÒNG, {n_cols} CỘT)"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Total Records: {n_rows} | Columns: {n_cols}"
        ws["A2"].font = Font(italic=True, color="595959")
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[4].height = 45
        ws.freeze_panes = "C5"
        for j, col in enumerate(df.columns, 1):
            letter = ws.cell(4, j).column_letter
            ws.column_dimensions[letter].width = 18
            if col in METRIC_ORDER:
                fmt = "#,##0;(#,##0);-" if col.endswith("vnd_mn") else "0.00"
                for r in range(5, 5 + n_rows):
                    ws.cell(r, j).number_format = fmt


def save_outputs(master: pd.DataFrame, with_metrics: pd.DataFrame, outdir: str):
    os.makedirs(outdir, exist_ok=True)
    p1 = os.path.join(outdir, "master_hopnhat.xlsx")
    p2 = os.path.join(outdir, "master_hopnhat_with_metrics.xlsx")
    write_xlsx(master, p1)
    write_xlsx(with_metrics, p2)
    master.to_csv(os.path.join(outdir, "master_hopnhat.csv"), index=False, encoding="utf-8-sig")
    with_metrics.to_csv(os.path.join(outdir, "master_hopnhat_with_metrics.csv"),
                        index=False, encoding="utf-8-sig")
    return p1, p2


# ----------------------------------------------------------------------------
# 7. MAIN
# ----------------------------------------------------------------------------
def main():
    global _REPORT_PATH
    ap = argparse.ArgumentParser(description="Hop nhat Data Pack thanh file master va tinh 17 chi tieu")
    ap.add_argument("--data-dir", default="01_data", help="Thu muc chua cac file csv cua Data Pack")
    ap.add_argument("--output", default="output_master", help="Thu muc xuat ket qua")
    args = ap.parse_args()
    _REPORT_PATH = os.path.join(args.output, "bao_cao_kiem_tra.txt")

    master = build_master(args.data_dir)
    master = add_checks(master)
    master = order_columns(master)
    quality_summary(master)

    with_metrics = order_columns(add_metrics(master.copy()))
    p1, p2 = save_outputs(master, with_metrics, args.output)

    log("\n" + "=" * 70)
    log("HOAN TAT")
    log("=" * 70)
    log(f"Master (khong chi tieu)  : {os.path.abspath(p1)}  [{master.shape[0]} dong x {master.shape[1]} cot]")
    log(f"Master + 17 chi tieu     : {os.path.abspath(p2)}  [{with_metrics.shape[0]} dong x {with_metrics.shape[1]} cot]")
    log("Buoc tiep theo: py benchmark_nganh.py --input "
        + os.path.join(args.output, "master_hopnhat_with_metrics.xlsx"))
    save_report_on_error()


if __name__ == "__main__":
    main()
